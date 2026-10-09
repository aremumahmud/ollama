"""Plain-text extraction for note attachments (PDF, DOCX)."""

import io

import fitz
from docx import Document as DocxDocument


def extract_pdf_text(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        return "\n\n".join(page.get_text() for page in doc).strip()
    finally:
        doc.close()


def extract_pdf_pages(data: bytes) -> list[dict]:
    """Returns [{"page": 1, "text": "..."}, ...], 1-indexed, one entry per page."""
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        return [{"page": i + 1, "text": page.get_text().strip()} for i, page in enumerate(doc)]
    finally:
        doc.close()


def extract_docx_text(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs).strip()


def extract_document_text(data: bytes, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf_text(data)
    if lower.endswith(".docx"):
        return extract_docx_text(data)
    raise ValueError("Only .pdf and .docx files are supported")
